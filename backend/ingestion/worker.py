"""No Django, database, network or source execution. Invoked as a bounded subprocess."""

import base64
import csv
import io
import json
import os
import re
import sys
import zipfile


class ParseFailure(ValueError):
    pass


def require(condition, message):
    if not condition:
        raise ParseFailure(message)


def text_bytes(data):
    return data.decode("utf-16" if data.startswith((b"\xff\xfe", b"\xfe\xff")) else "utf-8-sig")


def table_rows(records, limits, header_row, sheet):
    headers, rows, total = None, [], 0
    for number, cells in enumerate(records, 1):
        require(number <= limits["rows"] + header_row, "Row limit exceeded.")
        require(len(cells) <= limits["columns"], "Column limit exceeded.")
        cells = ["" if value is None else str(value) for value in cells]
        require(not any("\x00" in c for c in cells), "Unsupported null character in cell.")
        require(all(len(c) <= limits["cell_chars"] for c in cells), "Cell text limit exceeded.")
        total += sum(map(len, cells))
        require(total <= limits["text_chars"], "Extracted text limit exceeded.")
        if number < header_row:
            continue
        if number == header_row:
            headers = [c.strip() for c in cells]
            require(headers and all(headers), "Every column needs a nonempty header.")
            require(len(set(h.casefold() for h in headers)) == len(headers), "Duplicate headers.")
            require(all(len(h) <= 80 for h in headers), "Header text limit exceeded.")
            require(not any(h.startswith(("=", "+", "-", "@")) for h in headers), "Unsafe header.")
        elif any(c.strip() for c in cells):
            row_error = []
            if len(cells) != len(headers):
                row_error.append("Row width differs from headers.")
            rows.append(
                {
                    "number": number,
                    "location": f"{sheet}:row {number}",
                    "cells": dict(zip(headers, cells, strict=False)),
                    "errors": row_error,
                }
            )
    require(headers is not None and rows, "No data rows under the selected header.")
    return {"headers": headers, "rows": rows}


def csv_adapter(data, options, limits):
    source = text_bytes(data)
    delimiter = options.get("delimiter", "")
    if not delimiter:
        try:
            delimiter = csv.Sniffer().sniff(source[:8192], delimiters=",;\t|").delimiter
        except csv.Error as error:
            raise ParseFailure("Delimiter is ambiguous; select it explicitly.") from error
    csv.field_size_limit(limits["cell_chars"])
    result = table_rows(
        csv.reader(io.StringIO(source, newline=""), delimiter=delimiter, strict=True),
        limits,
        options.get("header_row", 1),
        "CSV",
    )
    return {**result, "sheets": [], "delimiter": delimiter}


def xlsx_adapter(data, options, limits):
    from defusedxml.ElementTree import fromstring
    from openpyxl import load_workbook

    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        entries = archive.infolist()
        require(len(entries) <= limits["archive_entries"], "Workbook archive entry limit exceeded.")
        require(len({e.filename for e in entries}) == len(entries), "Duplicate archive entries.")
        require(
            sum(e.file_size for e in entries) <= limits["unpacked_bytes"],
            "Expanded workbook limit exceeded.",
        )
        names = {e.filename for e in entries}
        require({"[Content_Types].xml", "xl/workbook.xml"} <= names, "Not an XLSX workbook.")
        for entry in entries:
            name = entry.filename.lower()
            require(not entry.flag_bits & 1, "Encrypted workbooks are unsupported.")
            require(
                not (name.startswith("/") or ".." in name.split("/") or "\\" in name),
                "Unsafe archive entry.",
            )
            require(
                entry.file_size <= max(1, entry.compress_size) * limits["compression_ratio"],
                "Workbook compression ratio limit exceeded.",
            )
            require(
                not any(
                    marker in name
                    for marker in ("vbaproject", "externallinks", "embeddings/", "activex")
                ),
                "Macros, external links and embedded objects are unsupported.",
            )
            if name.endswith((".xml", ".rels")):
                xml = archive.read(entry)
                root = fromstring(xml)
                require(
                    b"macroenabled" not in xml.lower(), "Macro-enabled workbooks are unsupported."
                )
                if name.endswith(".rels"):
                    require(
                        not any(node.attrib.get("TargetMode") == "External" for node in root),
                        "External workbook relationships are unsupported.",
                    )
    book = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
    try:
        require(len(book.sheetnames) <= limits["sheets"], "Sheet limit exceeded.")
        selected = options.get("sheet") or book.sheetnames[0]
        require(selected in book.sheetnames, "Selected sheet does not exist.")
        sheet = book[selected]
        # Ignore attacker-supplied dimensions; iterate actual XML cells under explicit caps.
        sheet.reset_dimensions()

        def records():
            for row in sheet.iter_rows():
                yield [
                    "[UNSUPPORTED FORMULA]" if cell.data_type == "f" else cell.value for cell in row
                ]

        result = table_rows(records(), limits, options.get("header_row", 1), selected)
        return {**result, "sheets": book.sheetnames, "sheet": selected}
    finally:
        book.close()


DOCUMENT_FIELDS = {
    "sku": "sku",
    "product name": "name",
    "name": "name",
    "manufacturer": "manufacturer",
    "brand": "brand",
    "description": "description",
    "package name": "package_name",
    "package type": "package_type",
    "material": "material",
    "units per package": "quantity",
    "length": "length",
    "width": "width",
    "height": "height",
    "weight": "mass",
    "mass": "mass",
}


def document_adapter(data, options, limits, format):
    if format == "pdf":
        from pypdf import PdfReader, filters

        filters.ZLIB_MAX_OUTPUT_LENGTH = limits["unpacked_bytes"]
        reader = PdfReader(io.BytesIO(data), strict=True)
        require(not reader.is_encrypted, "Encrypted PDF is unsupported.")
        require(len(reader.pages) <= limits["pages"], "Page limit exceeded.")
        # Active/attached content is never used; reject it rather than retain a disguised file.
        require(
            not reader.trailer["/Root"].get("/OpenAction") and not reader.attachments,
            "Active or attached PDF content is unsupported.",
        )
        pages = [page.extract_text() or "" for page in reader.pages]
    else:
        pages = [text_bytes(data)]
    require(sum(map(len, pages)) <= limits["text_chars"], "Document text limit exceeded.")
    headers = list(dict.fromkeys(DOCUMENT_FIELDS.values())) + ["dimension_unit", "mass_unit"]
    rows = []
    for number, page in enumerate(pages, 1):
        require("\x00" not in page, "Unsupported null character in document.")
        values, errors = {}, []
        for line in page.splitlines():
            match = re.fullmatch(r"\s*([A-Za-z ]{1,32})\s*:\s*(.{1,2000})\s*", line)
            if not match:
                continue
            field = DOCUMENT_FIELDS.get(match[1].strip().lower())
            if not field:
                continue
            value = match[2].strip()
            if field in values:
                errors.append(f"Multiple {field} labels on one page; split this document.")
            if field in {"length", "width", "height", "mass"}:
                physical = re.fullmatch(r"([0-9]+(?:\.[0-9]+)?)\s*(mm|cm|m|kg|g)", value)
                if physical:
                    value, unit = physical.groups()
                    unit_field = "mass_unit" if field == "mass" else "dimension_unit"
                    if unit_field in values and values[unit_field] != unit:
                        errors.append("Mixed dimension units on one page; correct explicitly.")
                    values[unit_field] = unit
            values[field] = value
        if values:
            rows.append(
                {"number": number, "location": f"page {number}", "cells": values, "errors": errors}
            )
    require(
        rows, "No supported labeled text found. Scans/OCR and arbitrary layouts are unsupported."
    )
    return {"headers": headers, "rows": rows, "sheets": []}


def main():
    # Hard limits on Linux production/CI. Windows still has input/output caps and wall timeout.
    if os.name == "posix":
        import resource

        resource.setrlimit(resource.RLIMIT_AS, (512 * 1024**2, 512 * 1024**2))
        resource.setrlimit(resource.RLIMIT_CPU, (10, 10))
    payload = json.load(sys.stdin)
    data = base64.b64decode(payload["data"], validate=True)
    format, options, limits = payload["format"], payload["options"], payload["limits"]
    try:
        if format == "csv":
            result = csv_adapter(data, options, limits)
        elif format == "xlsx":
            result = xlsx_adapter(data, options, limits)
        else:
            result = document_adapter(data, options, limits, format)
        print(json.dumps({"ok": True, **result}))
    except ParseFailure as error:
        print(json.dumps({"ok": False, "error": str(error)}))
    except Exception:
        print(
            json.dumps(
                {
                    "ok": False,
                    "error": "Malformed or unsupported file; check format, encoding and content.",
                }
            )
        )


if __name__ == "__main__":
    main()
